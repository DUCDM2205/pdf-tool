"""Conversions with explicit fidelity modes and local-only signing."""
import base64, io, importlib.util, os, shutil, subprocess, tempfile
from pathlib import Path
import fitz
from pdf_engine import InputError


def office_executable():
    candidates=[os.environ.get('FRIDAY_SOFFICE',''),shutil.which('soffice'),shutil.which('libreoffice'),
                'C:/Program Files/LibreOffice/program/soffice.exe','C:/Program Files (x86)/LibreOffice/program/soffice.exe']
    return next((str(p) for p in candidates if p and Path(p).is_file()),None)


def capabilities():
    return {'office':bool(office_executable()),'ocr':bool(shutil.which('tesseract') or Path('C:/Program Files/Tesseract-OCR/tesseract.exe').exists()),
            'word':bool(importlib.util.find_spec('pdf2docx')),'excel':bool(importlib.util.find_spec('openpyxl')),
            'powerpoint':bool(importlib.util.find_spec('pptx')),'certificate':bool(importlib.util.find_spec('pyhanko'))}


def office_to_pdf(name,data):
    exe=office_executable()
    if not exe:raise InputError('Chưa cài LibreOffice. Xem INSTALL.md để bật Word/Excel/PowerPoint → PDF.')
    suffix=Path(name).suffix.lower()
    if suffix not in {'.doc','.docx','.xls','.xlsx','.ppt','.pptx','.odt','.ods','.odp'}:raise InputError('Định dạng Office không hỗ trợ.')
    with tempfile.TemporaryDirectory(prefix='friday-office-') as temp:
        root=Path(temp);source=root/('input'+suffix);source.write_bytes(data)
        profile=(root/'profile').as_uri()
        result=subprocess.run([exe,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',str(root),str(source)],capture_output=True,timeout=180)
        output=root/'input.pdf'
        if result.returncode or not output.exists():raise InputError('LibreOffice không chuyển được file. Kiểm tra file có mật khẩu hoặc đang bị lỗi.')
        return output.read_bytes()


def convert_pdf(data,target,mode='editable'):
    if target=='docx':
        try:from pdf2docx import Converter
        except ImportError as exc:raise InputError('Thiếu pdf2docx. Chạy pip install -r requirements.txt theo INSTALL.md.') from exc
        with tempfile.TemporaryDirectory(prefix='friday-word-') as temp:
            src=Path(temp)/'input.pdf';out=Path(temp)/'output.docx';src.write_bytes(data)
            converter=Converter(str(src))
            try:converter.convert(str(out),multi_processing=False)
            finally:converter.close()
            return out.read_bytes()
    if target=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb=Workbook();wb.remove(wb.active)
        with fitz.open(stream=data,filetype='pdf') as doc:
            for i,page in enumerate(doc):
                ws=wb.create_sheet(f'Trang {i+1}')
                tables=page.find_tables().tables
                if tables:
                    for table in tables:
                        for row in table.extract():
                            # Keep PDF data as literal strings; never execute formulas.
                            ws.append([str(v or '') for v in row])
                            for cell in ws[ws.max_row]:cell.data_type='s'
                        ws.append([])
                else:
                    ws.append(['Không nhận diện được bảng; dưới đây là văn bản theo dòng.'])
                    for line in page.get_text().splitlines():
                        ws.append([line]);ws.cell(ws.max_row,1).data_type='s'
                for cell in ws[1]:cell.font=Font(bold=True)
                for col in ws.columns:
                    ws.column_dimensions[col[0].column_letter].width=min(65,max(12,max(len(str(c.value or '')) for c in col)+2))
        out=io.BytesIO();wb.save(out);return out.getvalue()
    if target=='pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs=Presentation()
        with fitz.open(stream=data,filetype='pdf') as doc:
            prs.slide_width=Inches(doc[0].rect.width/72);prs.slide_height=Inches(doc[0].rect.height/72)
            for page in doc:
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                sx=prs.slide_width/page.rect.width;sy=prs.slide_height/page.rect.height
                if mode=='visual':
                    pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
                    slide.shapes.add_picture(io.BytesIO(pix.tobytes('png')),0,0,width=prs.slide_width,height=prs.slide_height)
                else:
                    from editor_engine import blocks
                    for b in blocks(page):
                        x0,y0,x1,y1=b['bbox']
                        box=slide.shapes.add_textbox(int(x0*sx),int(y0*sy),max(1,int((x1-x0)*sx)),max(1,int((y1-y0)*sy)))
                        tf=box.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                        tf.word_wrap=True
                        for j,line in enumerate(b['text'].splitlines()):
                            p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line
                            p.font.size=Pt(b['size']);p.font.bold=b['bold'];p.font.italic=b['italic']
                    for info in page.get_image_info(xrefs=True):
                        if not info['xref']:continue
                        raw=doc.extract_image(info['xref'])['image'];x0,y0,x1,y1=info['bbox']
                        slide.shapes.add_picture(io.BytesIO(raw),int(x0*sx),int(y0*sy),width=max(1,int((x1-x0)*sx)),height=max(1,int((y1-y0)*sy)))
        out=io.BytesIO();prs.save(out);return out.getvalue()
    raise InputError('Định dạng xuất không hợp lệ.')


def sign_pdf(data,certificate,password):
    try:
        from pyhanko.sign import signers
        from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    except ImportError as exc:raise InputError('Thiếu pyHanko. Chạy cài requirements.txt trước khi ký số.') from exc
    raw=base64.b64decode(certificate.split(',')[-1],validate=True)
    if len(raw)>5*1024*1024:raise InputError('Chứng thư tối đa 5 MB.')
    try:
        signer=signers.SimpleSigner.load_pkcs12_data(raw,other_certs=[],passphrase=password.encode() if password else None)
        if signer is None:raise ValueError('Không đọc được chứng thư')
        output=io.BytesIO()
        signers.sign_pdf(IncrementalPdfFileWriter(io.BytesIO(data)),signers.PdfSignatureMetadata(field_name='FridayDigitalSignature'),signer=signer,output=output)
        return output.getvalue()
    except Exception as exc:raise InputError('Ký số thất bại. Kiểm tra file P12/PFX chứa private key và mật khẩu.') from exc
