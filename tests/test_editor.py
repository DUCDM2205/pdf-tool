import io, unittest
import fitz
from PIL import Image
from pypdf import PdfReader
from pdf_engine import Source, InputError, build_pdf
from editor_engine import scan, edit
from converters import convert_pdf


def fixture():
    d=fitz.open();p=d.new_page(width=400,height=500)
    p.insert_text((25,40),'EDIT THIS',fontsize=12)
    p.insert_text((25,200),'KEEP NEIGHBOUR',fontsize=12)
    out=io.BytesIO();Image.new('RGB',(80,50),'red').save(out,'PNG')
    p.insert_image(fitz.Rect(240,250,320,300),stream=out.getvalue())
    p.insert_text((25,350),'SECRET-987654',fontsize=12)
    return d.tobytes()


class EditorTest(unittest.TestCase):
    def setUp(self):
        self.sources={'s':Source('test.pdf',fixture(),'pdf')}
        self.entry={'source':'s','index':0,'ops':[],'crop':{}}

    def test_unicode_wrap_preserves_neighbour_and_rejects_overflow(self):
        item=scan(self.entry,self.sources)['text'][0]
        op={'kind':'text','id':item['id'],'bbox':[25,25,215,130],
            'text':'Nội dung tiếng Việt được thay đổi và tự động xuống dòng.','size':12}
        data=edit(self.entry,self.sources,op)
        with fitz.open(stream=data,filetype='pdf') as d:
            text=d[0].get_text()
            self.assertNotIn('EDIT THIS',text);self.assertIn('KEEP NEIGHBOUR',text)
            self.assertIn('tiếng Việt',text);self.assertIn('xuống dòng.',text)
            self.assertEqual(len(d[0].get_images()),1)
        with self.assertRaises(InputError):edit(self.entry,self.sources,{**op,'text':'Too much '*500})
        with self.assertRaises(InputError):edit(self.entry,self.sources,{**op,'bbox':[25,25,220,220]})

    def test_text_over_background_preserves_image(self):
        d=fitz.open();p=d.new_page(width=400,height=500)
        buf=io.BytesIO();Image.new('RGB',(400,500),(210,225,240)).save(buf,'PNG')
        p.insert_image(p.rect,stream=buf.getvalue())
        p.insert_text((25,50),'ORIGINAL TEXT',fontsize=12)
        p.insert_text((25,220),'NEIGHBOUR',fontsize=12)
        self.sources['s']=Source('background.pdf',d.tobytes(),'pdf')
        item=scan(self.entry,self.sources)['text'][0]
        op={'kind':'text','id':item['id'],'bbox':[25,30,220,130],'text':'Updated text over background','size':12}
        output=edit(self.entry,self.sources,op)
        with fitz.open(stream=output,filetype='pdf') as result:
            self.assertIn('Updated text',result[0].get_text())
            self.assertNotIn('ORIGINAL',result[0].get_text())
            self.assertIn('NEIGHBOUR',result[0].get_text())
            self.assertEqual(result[0].get_pixmap().pixel(30,100),(210,225,240))
        with self.assertRaises(InputError):
            edit(self.entry,self.sources,{**op,'bbox':[25,30,220,240]})

    def test_new_image_overlap_requires_explicit_choice(self):
        op={'kind':'addText','bbox':[240,250,320,295],'text':'Caption','size':10}
        with self.assertRaises(InputError):edit(self.entry,self.sources,op)
        output=edit(self.entry,self.sources,{**op,'allowImageOverlap':True})
        with fitz.open(stream=output,filetype='pdf') as result:
            self.assertIn('Caption',result[0].get_text())
            self.assertEqual(result[0].get_pixmap().pixel(250,290),(255,0,0))

    def test_redaction_removes_text_from_serialized_pdf(self):
        data=edit(self.entry,self.sources,{'kind':'redact','bbox':[20,330,190,360]})
        with fitz.open(stream=data,filetype='pdf') as d:
            self.assertNotIn('SECRET',d[0].get_text())
            self.assertIn('KEEP NEIGHBOUR',d[0].get_text())
            streams=b''.join(d.xref_stream(i) or b'' for i in range(1,d.xref_length()) if d.xref_is_stream(i))
            self.assertNotIn(b'SECRET',streams)
        self.assertNotIn('SECRET',PdfReader(io.BytesIO(data)).pages[0].extract_text())

    def test_image_move_keeps_text(self):
        data=edit(self.entry,self.sources,{'kind':'image','id':'0','bbox':[240,390,360,460]})
        with fitz.open(stream=data,filetype='pdf') as d:
            self.assertIn('KEEP NEIGHBOUR',d[0].get_text())
            pix=d[0].get_pixmap()
            self.assertEqual(pix.pixel(260,270),(255,255,255))
            self.assertEqual(pix.pixel(270,410),(255,0,0))

    def test_form_tree_survives_export(self):
        data=edit(self.entry,self.sources,{'kind':'form','bbox':[25,420,180,450],'text':'Alice'})
        sources={'f':Source('form.pdf',data,'pdf')}
        exported=build_pdf([{'source':'f','index':0,'ops':[],'crop':{}}],sources)
        reader=PdfReader(io.BytesIO(exported))
        fields=reader.get_fields()
        self.assertTrue(fields)
        self.assertIn('Alice',[f.get('/V') for f in fields.values()])
        self.assertTrue(reader.pages[0]['/Annots'][0].get_object()['/AP']['/N'])

    def test_rotated_page_scan_and_edit_match(self):
        entry={**self.entry,'ops':['right','horizontal']}
        scanned=scan(entry,self.sources)
        self.assertEqual(scanned['width'],500)
        data=edit(entry,self.sources,{'kind':'addText','bbox':[330,200,490,245],'text':'Added','size':12})
        with fitz.open(stream=data,filetype='pdf') as d:self.assertIn('Added',d[0].get_text())

    def test_table_export_xlsx_pptx(self):
        data=edit(self.entry,self.sources,{'kind':'table','bbox':[20,70,220,165],'rows':[['Name','Value'],['Item','123']],'size':11})
        with fitz.open(stream=data,filetype='pdf') as d:
            self.assertEqual(len(d[0].find_tables().tables),1)
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(convert_pdf(data,'xlsx')))
        self.assertEqual(wb.active['B2'].value,'123')
        from pptx import Presentation
        for mode in ['editable','visual']:
            prs=Presentation(io.BytesIO(convert_pdf(data,'pptx',mode)))
            self.assertEqual(len(prs.slides),1)

if __name__=='__main__':unittest.main()
